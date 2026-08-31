import csv
import json
import subprocess
import sys
import unittest
from datetime import datetime
from pathlib import Path
from subprocess import run
from tempfile import TemporaryDirectory
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from PIL import Image
from reportlab.pdfgen import canvas

sys.path.insert(0, str(Path(__file__).parents[1]))

from benchmarks.verifiers import verify_case


class BenchmarkVerifierTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = TemporaryDirectory()
        self.home = Path(self.temp_dir.name)
        self.fixture = self.home / "fixtures" / "LD01"
        self.run = self.home / "runs" / "LD01"
        self.fixture.mkdir(parents=True)
        self.run.mkdir(parents=True)

    def tearDown(self):
        self.temp_dir.cleanup()

    def _case(self, verification: dict) -> dict:
        return {
            "id": "LD01",
            "fixture_root": "fixtures/LD01",
            "run_root": "runs/LD01",
            "verification": verification,
        }

    def test_text_supports_expected_file_and_contains_rule(self):
        (self.fixture / "expected.txt").write_text("finished", encoding="utf-8")
        (self.run / "result.txt").write_text("finished", encoding="utf-8")
        exact = self._case(
            {"kind": "text", "output": "result.txt", "expected": "expected.txt"}
        )
        contains = self._case(
            {"kind": "text", "output": "result.txt", "contains": "finish"}
        )

        self.assertTrue(verify_case(exact, self.home)["success"])
        self.assertTrue(verify_case(contains, self.home)["success"])

    def test_same_file_compares_output_with_original_fixture(self):
        (self.fixture / "source.txt").write_bytes(b"same bytes")
        (self.run / "copy.txt").write_bytes(b"same bytes")
        case = self._case(
            {"kind": "same_file", "output": "copy.txt", "expected": "source.txt"}
        )

        self.assertTrue(verify_case(case, self.home)["success"])

    def test_same_file_uses_upstream_text_comparison(self):
        (self.fixture / "source.txt").write_bytes(b"line one\r\nline two\r\n")
        (self.run / "copy.txt").write_bytes(b"line one\nline two\n")
        case = self._case(
            {"kind": "same_file", "output": "copy.txt", "expected": "source.txt"}
        )

        self.assertTrue(verify_case(case, self.home)["success"])

    def test_all_files_moved_uses_fixture_as_expected_member_list(self):
        (self.fixture / "Documents").mkdir()
        (self.fixture / "Documents" / "a.docx").touch()
        (self.fixture / "Documents" / "b.docx").touch()
        (self.run / "Documents" / "Archive").mkdir(parents=True)
        (self.run / "Documents" / "Archive" / "a.docx").touch()
        (self.run / "Documents" / "Archive" / "b.docx").touch()
        case = self._case(
            {
                "kind": "all_files_moved",
                "source": "Documents",
                "pattern": "*.docx",
                "destination": "Documents/Archive",
            }
        )

        self.assertTrue(verify_case(case, self.home)["success"])

    def test_all_files_moved_rejects_non_docx_file_in_archive(self):
        (self.fixture / "Documents").mkdir()
        (self.fixture / "Documents" / "a.docx").touch()
        archive = self.run / "Documents" / "Archive"
        archive.mkdir(parents=True)
        (archive / "a.docx").touch()
        (archive / "notes.txt").touch()
        case = self._case(
            {
                "kind": "all_files_moved",
                "source": "Documents",
                "pattern": "*.docx",
                "destination": "Documents/Archive",
            }
        )

        self.assertFalse(verify_case(case, self.home)["success"])

    def test_image_size_reads_real_output_dimensions(self):
        output = self.run / "canvas.png"
        Image.new("RGB", (800, 600)).save(output)
        case = self._case(
            {"kind": "image_size", "output": "canvas.png", "width": 800, "height": 600}
        )

        self.assertTrue(verify_case(case, self.home)["success"])

    def test_docx_fonts_checks_each_text_run(self):
        output = self.run / "document.docx"
        document = Document()
        run = document.add_paragraph().add_run("formatted")
        run.font.name = "Times New Roman"
        document.save(output)
        case = self._case(
            {"kind": "docx_fonts", "output": "document.docx", "font": "Times New Roman"}
        )

        self.assertTrue(verify_case(case, self.home)["success"])

    def test_docx_fonts_includes_blank_runs_like_upstream(self):
        output = self.run / "document.docx"
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("formatted").font.name = "Times New Roman"
        paragraph.add_run(" ").font.name = "Arial"
        document.save(output)
        case = self._case(
            {"kind": "docx_fonts", "output": "document.docx", "font": "Times New Roman"}
        )

        self.assertFalse(verify_case(case, self.home)["success"])

    def test_docx_fonts_accepts_empty_document_like_upstream(self):
        output = self.run / "document.docx"
        Document().save(output)
        case = self._case(
            {"kind": "docx_fonts", "output": "document.docx", "font": "Times New Roman"}
        )

        self.assertTrue(verify_case(case, self.home)["success"])

    def test_png_names_compares_trimmed_line_set(self):
        fixture_pictures = self.fixture / "Pictures"
        fixture_pictures.mkdir()
        (fixture_pictures / "a.png").touch()
        (fixture_pictures / "b.png").touch()
        pictures = self.run / "Pictures"
        pictures.mkdir()
        (pictures / "a.png").touch()
        (pictures / "b.png").touch()
        (pictures / "png_files.txt").write_text("b.png\na.png\n", encoding="utf-8")
        case = self._case(
            {
                "kind": "png_names",
                "output": "Pictures/png_files.txt",
                "directory": "Pictures",
                "pattern": "*.png",
            }
        )

        self.assertTrue(verify_case(case, self.home)["success"])

    def test_png_names_compares_the_run_directory_not_static_fixture(self):
        fixture_pictures = self.fixture / "Pictures"
        fixture_pictures.mkdir()
        (fixture_pictures / "a.png").touch()
        pictures = self.run / "Pictures"
        pictures.mkdir()
        (pictures / "a.png").touch()
        (pictures / "new.PNG").touch()
        (pictures / "png_files.txt").write_text("a.png\nnew.PNG\n", encoding="utf-8")
        case = self._case(
            {
                "kind": "png_names",
                "output": "Pictures/png_files.txt",
                "directory": "Pictures",
                "pattern": "*.png",
                "expected_members": ["a.png"],
            }
        )

        self.assertTrue(verify_case(case, self.home)["success"])

    def test_png_names_preserves_line_content_like_upstream(self):
        pictures = self.run / "Pictures"
        pictures.mkdir()
        (pictures / "a.png").touch()
        (pictures / "png_files.txt").write_text("a.png \n", encoding="utf-8")
        case = self._case(
            {
                "kind": "png_names",
                "output": "Pictures/png_files.txt",
                "directory": "Pictures",
                "pattern": "*.png",
            }
        )

        self.assertFalse(verify_case(case, self.home)["success"])

    def test_missing_output_is_a_normal_failed_evaluation(self):
        case = self._case(
            {"kind": "text", "output": "missing.txt", "contains": "done"}
        )

        result = verify_case(case, self.home)

        self.assertFalse(result["success"])
        self.assertIn("missing", result["detail"])

    def test_hidden_file_uses_windows_file_attribute(self):
        output = self.run / "secret.txt"
        output.write_text("private", encoding="utf-8")
        run(["attrib", "+H", str(output)], check=True)

        result = verify_case(
            self._case({"kind": "hidden", "output": "secret.txt"}), self.home
        )

        self.assertTrue(result["success"])

    def test_workbook_values_compares_real_workbooks(self):
        for path in (self.fixture / "gold.xlsx", self.run / "result.xlsx"):
            workbook = Workbook()
            workbook.active.append(["name", "score"])
            workbook.active.append(["Ada", 10])
            workbook.save(path)

        result = verify_case(
            self._case(
                {
                    "kind": "workbook_values",
                    "output": "result.xlsx",
                    "expected": "gold.xlsx",
                    "rules": [
                        {"type": "sheet_data", "sheet_idx0": 0, "sheet_idx1": "EI0"}
                    ],
                }
            ),
            self.home,
        )

        self.assertTrue(result["success"])

    def test_workbook_values_uses_declared_sheet_references(self):
        actual = Workbook()
        actual.active.title = "decoy"
        actual.active.append(["unchanged"])
        actual.create_sheet("result").append(["wrong"])
        actual.save(self.run / "result.xlsx")

        expected = Workbook()
        expected.active.title = "decoy"
        expected.active.append(["unchanged"])
        expected.create_sheet("gold").append(["correct"])
        expected.save(self.fixture / "gold.xlsx")

        case = self._case(
            {
                "kind": "workbook_values",
                "output": "result.xlsx",
                "expected": "gold.xlsx",
                "rules": [
                    {"type": "sheet_data", "sheet_idx0": 1, "sheet_idx1": "EI1"}
                ],
            }
        )

        self.assertFalse(verify_case(case, self.home)["success"])

        actual = Workbook()
        actual.active.title = "decoy"
        actual.active.append(["different"])
        actual.create_sheet("result").append(["correct"])
        actual.save(self.run / "result.xlsx")

        self.assertTrue(verify_case(case, self.home)["success"])

    def test_workbook_values_uses_declared_numeric_precision(self):
        actual = Workbook()
        actual.active.append(["score"])
        actual.active.append([1.23456])
        actual.save(self.run / "result.xlsx")
        expected = Workbook()
        expected.active.append(["score"])
        expected.active.append([1.234559])
        expected.save(self.fixture / "gold.xlsx")
        case = self._case(
            {
                "kind": "workbook_values",
                "output": "result.xlsx",
                "expected": "gold.xlsx",
                "rules": [
                    {
                        "type": "sheet_data",
                        "sheet_idx0": 0,
                        "sheet_idx1": "EI0",
                        "precision": 4,
                    }
                ],
            }
        )

        self.assertTrue(verify_case(case, self.home)["success"])

    def test_workbook_values_preserves_upstream_dataframe_types(self):
        actual = Workbook()
        actual.active.append(["value"])
        actual.active.append([True])
        actual.save(self.run / "result.xlsx")
        expected = Workbook()
        expected.active.append(["value"])
        expected.active.append([1])
        expected.save(self.fixture / "gold.xlsx")
        case = self._case(
            {
                "kind": "workbook_values",
                "output": "result.xlsx",
                "expected": "gold.xlsx",
                "rules": [
                    {"type": "sheet_data", "sheet_idx0": 0, "sheet_idx1": "EI0"}
                ],
            }
        )

        self.assertFalse(verify_case(case, self.home)["success"])

    def test_pdf_text_compares_extracted_text(self):
        for path in (self.fixture / "gold.pdf", self.run / "result.pdf"):
            document = canvas.Canvas(str(path))
            document.drawString(72, 720, "OmniAct benchmark")
            document.save()

        result = verify_case(
            self._case(
                {"kind": "pdf_text", "output": "result.pdf", "expected": "gold.pdf"}
            ),
            self.home,
        )

        self.assertTrue(result["success"])

    def test_pdf_text_keeps_upstream_partial_score(self):
        actual = canvas.Canvas(str(self.run / "result.pdf"))
        actual.drawString(72, 720, "OmniAct benchmark result")
        actual.save()
        expected = canvas.Canvas(str(self.fixture / "gold.pdf"))
        expected.drawString(72, 720, "OmniAct benchmark")
        expected.save()

        result = verify_case(
            self._case(
                {"kind": "pdf_text", "output": "result.pdf", "expected": "gold.pdf"}
            ),
            self.home,
        )

        self.assertFalse(result["success"])
        self.assertGreater(result["score"], 0.0)
        self.assertLess(result["score"], 1.0)

    @patch("benchmarks.verifiers.subprocess.run")
    @patch("benchmarks.verifiers.which", return_value="7z.exe")
    def test_seven_zip_rejects_archive_that_does_not_require_password(
        self, _which, process
    ):
        output = self.run / "OldProjects.7z"
        output.touch()
        process.side_effect = [
            subprocess.CompletedProcess([], 0, stdout="Everything is Ok", stderr=""),
            subprocess.CompletedProcess([], 0, stdout="Everything is Ok", stderr=""),
        ]
        case = self._case(
            {
                "kind": "seven_zip",
                "output": "OldProjects.7z",
                "password": "12345",
                "expected_members": ["OldProjects/example.txt"],
            }
        )

        self.assertFalse(verify_case(case, self.home)["success"])

    @patch("benchmarks.verifiers.subprocess.run")
    @patch("benchmarks.verifiers.which", return_value="7z.exe")
    def test_seven_zip_accepts_encrypted_archive_with_expected_member(
        self, _which, process
    ):
        output = self.run / "OldProjects.7z"
        output.touch()
        process.side_effect = [
            subprocess.CompletedProcess([], 0, stdout="Everything is Ok", stderr=""),
            subprocess.CompletedProcess([], 2, stdout="", stderr="Wrong password"),
            subprocess.CompletedProcess(
                [],
                0,
                stdout=(
                    f"Path = {output}\nType = 7z\n\n"
                    "Path = OldProjects\\example.txt\nSize = 7\n"
                ),
                stderr="",
            ),
        ]
        case = self._case(
            {
                "kind": "seven_zip",
                "output": "OldProjects.7z",
                "password": "12345",
                "expected_members": ["OldProjects/example.txt"],
            }
        )

        self.assertTrue(verify_case(case, self.home)["success"])

    def test_yaml_records_checks_nested_values(self):
        output = self.run / "_config.yml"
        output.write_text(
            "name: Test Account\nauthor:\n  name: Test Account\n  email: Test@gmail.com\n",
            encoding="utf-8",
        )
        case = self._case(
            {
                "kind": "yaml_records",
                "output": "_config.yml",
                "rules": [
                    {"path": "name", "equals": "Test Account"},
                    {"path": "author.email", "equals": "Test@gmail.com"},
                ],
            }
        )

        self.assertTrue(verify_case(case, self.home)["success"])

    def test_malformed_yaml_is_a_normal_failed_evaluation(self):
        (self.run / "_config.yml").write_text("author: [", encoding="utf-8")
        case = self._case(
            {
                "kind": "yaml_records",
                "output": "_config.yml",
                "rules": [{"path": "author.name", "equals": "Test Account"}],
            }
        )

        result = verify_case(case, self.home)

        self.assertFalse(result["success"])

    @patch("benchmarks.verifiers.psutil.Process")
    @patch("benchmarks.verifiers.psutil.pid_exists", return_value=True)
    @patch("benchmarks.verifiers._vscode_profile_in_use", return_value=True)
    def test_terminal_command_requires_transcript_and_open_workspace(
        self, _profile_in_use, _pid_exists, process
    ):
        (self.run / "project").mkdir()
        storage = self.run / ".vscode-profile" / "User" / "workspaceStorage" / "one"
        storage.mkdir(parents=True)
        (storage / "workspace.json").write_text(
            json.dumps({"folder": (self.run / "project").resolve().as_uri()}),
            encoding="utf-8",
        )
        (self.run / "terminal-transcript.txt").write_text("code .\n", encoding="utf-8")
        (self.run / "case-processes.json").write_text(
            json.dumps({"pids": [42], "terminal_pid": 42}), encoding="utf-8"
        )
        process.return_value.cwd.return_value = str(self.run / "project")
        case = self._case(
            {
                "kind": "terminal_command_and_workspace",
                "workspace": "project",
                "terminal_transcript": "terminal-transcript.txt",
                "command": "code .",
                "vscode_profile": ".vscode-profile",
            }
        )

        self.assertTrue(verify_case(case, self.home)["success"])

    @patch("benchmarks.verifiers.psutil.Process")
    @patch("benchmarks.verifiers.psutil.pid_exists", return_value=True)
    @patch("benchmarks.verifiers._vscode_profile_in_use", return_value=True)
    def test_terminal_command_rejects_same_named_workspace_at_other_path(
        self, _profile_in_use, _pid_exists, process
    ):
        (self.run / "project").mkdir()
        other = self.home / "other" / "project"
        other.mkdir(parents=True)
        storage = self.run / ".vscode-profile" / "User" / "workspaceStorage" / "one"
        storage.mkdir(parents=True)
        workspace_state = storage / "workspace.json"
        workspace_state.write_text(
            json.dumps({"folder": other.resolve().as_uri()}), encoding="utf-8"
        )
        (self.run / "terminal-transcript.txt").write_text("code .\n", encoding="utf-8")
        (self.run / "case-processes.json").write_text(
            json.dumps({"pids": [42], "terminal_pid": 42}), encoding="utf-8"
        )
        process.return_value.cwd.return_value = str(self.run / "project")
        case = self._case(
            {
                "kind": "terminal_command_and_workspace",
                "workspace": "project",
                "terminal_transcript": "terminal-transcript.txt",
                "command": "code .",
                "vscode_profile": ".vscode-profile",
            }
        )

        self.assertFalse(verify_case(case, self.home)["success"])
        workspace_state.write_text(
            json.dumps({"folder": (self.run / "project").resolve().as_uri()}),
            encoding="utf-8",
        )
        self.assertTrue(verify_case(case, self.home)["success"])

    @patch("benchmarks.verifiers._vscode_profile_in_use", return_value=False)
    def test_terminal_command_rejects_stale_workspace_state(self, _profile_in_use):
        (self.run / "project").mkdir()
        storage = self.run / ".vscode-profile" / "User" / "workspaceStorage" / "one"
        storage.mkdir(parents=True)
        (storage / "workspace.json").write_text(
            json.dumps({"folder": (self.run / "project").resolve().as_uri()}),
            encoding="utf-8",
        )
        (self.run / "terminal-transcript.txt").write_text("code .\n", encoding="utf-8")
        (self.run / "case-processes.json").write_text(
            json.dumps({"pids": [42], "terminal_pid": 42}), encoding="utf-8"
        )
        case = self._case(
            {
                "kind": "terminal_command_and_workspace",
                "workspace": "project",
                "terminal_transcript": "terminal-transcript.txt",
                "command": "code .",
                "vscode_profile": ".vscode-profile",
            }
        )

        self.assertFalse(verify_case(case, self.home)["success"])

    def test_process_and_transcript_requires_recorded_process_to_exit(self):
        process = run([sys.executable, "-c", "pass"], check=True)
        metadata = {"pids": [process.returncode + 999999]}
        (self.run / "case-writer-process.json").write_text(
            json.dumps(metadata), encoding="utf-8"
        )
        (self.run / "terminal-transcript.txt").write_text(
            "Stop-Process -Id 999999", encoding="utf-8"
        )
        case = self._case(
            {
                "kind": "process_and_transcript",
                "process_metadata": "case-writer-process.json",
                "terminal_transcript": "terminal-transcript.txt",
                "required_terminal_evidence": "Stop-Process",
            }
        )

        self.assertTrue(verify_case(case, self.home)["success"])

    def test_process_and_transcript_rejects_stop_process_for_unrelated_pid(self):
        (self.run / "case-writer-process.json").write_text(
            json.dumps({"pids": [999999]}), encoding="utf-8"
        )
        (self.run / "terminal-transcript.txt").write_text(
            "Stop-Process -Id 123456", encoding="utf-8"
        )
        case = self._case(
            {
                "kind": "process_and_transcript",
                "process_metadata": "case-writer-process.json",
                "terminal_transcript": "terminal-transcript.txt",
                "required_terminal_evidence": "Stop-Process",
            }
        )

        self.assertFalse(verify_case(case, self.home)["success"])

    def test_process_and_transcript_accepts_process_name_termination(self):
        (self.run / "case-writer-process.json").write_text(
            json.dumps({"pids": [999999]}), encoding="utf-8"
        )
        (self.run / "terminal-transcript.txt").write_text(
            "Stop-Process -Name soffice", encoding="utf-8"
        )
        case = self._case(
            {
                "kind": "process_and_transcript",
                "process_metadata": "case-writer-process.json",
                "terminal_transcript": "terminal-transcript.txt",
                "required_terminal_evidence": "Stop-Process",
            }
        )

        self.assertTrue(verify_case(case, self.home)["success"])

    def test_speedtest_export_checks_csv_timestamp_freshness(self):
        output_dir = self.run / "Test" / "Speed"
        output_dir.mkdir(parents=True)
        output = output_dir / "Speedtest Results Export-1.csv"
        with output.open("w", newline="", encoding="utf-8") as file:
            writer = csv.writer(file)
            writer.writerow(["Server", "Date"])
            writer.writerow(["Local", datetime.now().strftime("%Y-%m-%d %H:%M")])
        case = self._case(
            {
                "kind": "speedtest_export_time",
                "output_directory": "Test/Speed",
                "file_glob": "Speedtest Results Export-*.csv",
                "timestamp_row": 2,
                "timestamp_column": 2,
                "max_age_minutes": 60,
            }
        )

        self.assertTrue(verify_case(case, self.home)["success"])

    def test_speedtest_export_uses_latest_matching_export(self):
        output_dir = self.run / "Test" / "Speed"
        output_dir.mkdir(parents=True)
        old = output_dir / "Speedtest Results Export-1.csv"
        new = output_dir / "Speedtest Results Export-2.csv"
        for path, value in ((old, "2000-01-01 00:00"), (new, datetime.now().strftime("%Y-%m-%d %H:%M"))):
            with path.open("w", newline="", encoding="utf-8") as file:
                writer = csv.writer(file)
                writer.writerow(["Server", "Date"])
                writer.writerow(["Local", value])
        old.touch()
        new.touch()
        case = self._case(
            {
                "kind": "speedtest_export_time",
                "output_directory": "Test/Speed",
                "file_glob": "Speedtest Results Export-*.csv",
                "timestamp_row": 2,
                "timestamp_column": 2,
                "max_age_minutes": 60,
            }
        )

        self.assertTrue(verify_case(case, self.home)["success"])

    def test_short_speedtest_csv_is_a_normal_failed_evaluation(self):
        output_dir = self.run / "Test" / "Speed"
        output_dir.mkdir(parents=True)
        (output_dir / "Speedtest Results Export-1.csv").write_text(
            "Server,Date\n", encoding="utf-8"
        )
        case = self._case(
            {
                "kind": "speedtest_export_time",
                "output_directory": "Test/Speed",
                "file_glob": "Speedtest Results Export-*.csv",
                "timestamp_row": 2,
                "timestamp_column": 2,
                "max_age_minutes": 60,
            }
        )

        result = verify_case(case, self.home)

        self.assertFalse(result["success"])

    def test_structured_text_counts_timestamp_lines(self):
        lines = [f"12:00:{second:02d} CPU %user %nice %system %iowait %steal %idle" for second in range(31)]
        (self.run / "report.txt").write_text("\n".join(lines), encoding="utf-8")
        case = self._case(
            {
                "kind": "structured_text",
                "output": "report.txt",
                "required_text": ["CPU", "%user", "%idle"],
                "timestamp_line_count": 31,
            }
        )

        self.assertTrue(verify_case(case, self.home)["success"])


if __name__ == "__main__":
    unittest.main()
