import csv
import json
import os
import re
import stat
import subprocess
from datetime import datetime
from pathlib import Path
from shutil import which
from typing import Any, Callable
from urllib.parse import unquote, urlsplit

import pandas as pd
import pymupdf as fitz
import psutil
import yaml
from docx import Document
from PIL import Image
from pypdf import PdfReader
from rapidfuzz import fuzz


Check = Callable[[Path, Path, dict[str, Any]], tuple[bool | float, str]]


def verify_case(case: dict[str, Any], bench_home: Path | str) -> dict[str, Any]:
    fixture = Path(bench_home) / case["fixture_root"]
    run = Path(bench_home) / case["run_root"]
    verification = case["verification"]
    kind = verification["kind"]
    check = CHECKS.get(kind)
    if check is None:
        raise ValueError(f"{case['id']}: unsupported verification kind: {kind}")

    try:
        outcome, detail = check(run, fixture, verification)
        score = float(outcome)
    except (
        OSError,
        ValueError,
        KeyError,
        TypeError,
        IndexError,
        RuntimeError,
        yaml.YAMLError,
    ) as error:
        score, detail = 0.0, str(error)
    return {
        "case_id": case["id"],
        "success": score == 1.0,
        "score": score,
        "detail": detail,
    }


def _text(run: Path, fixture: Path, spec: dict[str, Any]) -> tuple[bool, str]:
    output = run / spec["output"]
    actual = output.read_text(encoding="utf-8")
    if "contains" in spec:
        success = spec["contains"] in actual
    else:
        success = actual == (fixture / spec["expected"]).read_text(encoding="utf-8")
    return success, str(output)


def _same_file(run: Path, fixture: Path, spec: dict[str, Any]) -> tuple[bool, str]:
    output = run / spec["output"]
    actual = output.read_text(encoding="utf-8")
    expected = (fixture / spec["expected"]).read_text(encoding="utf-8")
    success = actual == expected
    return success, str(output)


def _hidden(run: Path, fixture: Path, spec: dict[str, Any]) -> tuple[bool, str]:
    output = run / spec["output"]
    attributes = os.stat(output).st_file_attributes
    return bool(attributes & stat.FILE_ATTRIBUTE_HIDDEN), str(output)


def _all_files_moved(
    run: Path, fixture: Path, spec: dict[str, Any]
) -> tuple[bool, str]:
    source = Path(spec["source"])
    source_path = run / source
    archive = run / spec["destination"]
    if not archive.is_dir():
        return False, f"archive_missing={archive}"
    archive_files = [path for path in archive.iterdir() if path.is_file()]
    remaining = [
        path.name
        for path in source_path.iterdir()
        if path.is_file() and path.suffix.casefold() == ".docx"
    ]
    invalid = [
        path.name for path in archive_files if path.suffix.casefold() != ".docx"
    ]
    success = bool(archive_files) and not invalid and not remaining
    return success, f"archive_files={sorted(path.name for path in archive_files)}, remaining={remaining}"


def _image_size(run: Path, fixture: Path, spec: dict[str, Any]) -> tuple[bool, str]:
    output = run / spec["output"]
    with Image.open(output) as image:
        size = image.size
    expected = (spec["width"], spec["height"])
    return size == expected, f"actual={size}, expected={expected}"


def _docx_fonts(run: Path, fixture: Path, spec: dict[str, Any]) -> tuple[bool, str]:
    output = run / spec["output"]
    document = Document(output)
    runs = [
        item
        for paragraph in document.paragraphs
        for item in paragraph.runs
    ]
    success = all(item.font.name == spec["font"] for item in runs)
    return success, f"runs={len(runs)}, expected_font={spec['font']}"


def _workbook_values(
    run: Path, fixture: Path, spec: dict[str, Any]
) -> tuple[bool, str]:
    output = run / spec["output"]
    expected = fixture / spec["expected"]
    with pd.ExcelFile(output) as actual_book, pd.ExcelFile(expected) as expected_book:
        rules = spec["rules"]
        for index, rule in enumerate(rules):
            if rule["type"] != "sheet_data":
                raise ValueError(f"unsupported workbook rule: {rule['type']}")
            precision = rule.get("precision", 4)
            actual = _sheet_data_frame(
                rule["sheet_idx0"], actual_book, expected_book
            ).round(precision)
            gold = _sheet_data_frame(
                rule["sheet_idx1"], actual_book, expected_book
            ).round(precision)
            if not actual.equals(gold):
                return False, (
                    f"failed_rule={index}, actual_rows={len(actual)}, "
                    f"expected_rows={len(gold)}"
                )
    return True, f"checked_sheet_data_rules={len(rules)}"


def _sheet_data_frame(
    reference: Any, actual_book: pd.ExcelFile, expected_book: pd.ExcelFile
) -> pd.DataFrame:
    if isinstance(reference, int):
        return pd.read_excel(actual_book, actual_book.sheet_names[reference])
    if isinstance(reference, str) and reference.startswith("EI"):
        return pd.read_excel(expected_book, expected_book.sheet_names[int(reference[2:])])
    raise ValueError(f"unsupported sheet_data reference: {reference}")


def _png_names(run: Path, fixture: Path, spec: dict[str, Any]) -> tuple[bool, str]:
    output = run / spec["output"]
    actual = set(output.read_text(encoding="utf-8").splitlines())
    expected = {
        path.name
        for path in (run / spec["directory"]).iterdir()
        if path.is_file() and path.suffix.casefold() == ".png"
    }
    return actual == expected, f"actual={sorted(actual)}, expected={sorted(expected)}"


def _seven_zip(run: Path, fixture: Path, spec: dict[str, Any]) -> tuple[bool, str]:
    executable = which("7z")
    if not executable:
        raise RuntimeError("7z is not available")
    output = run / spec["output"]
    password = spec["password"]
    tested = subprocess.run(
        [executable, "t", f"-p{password}", "-y", str(output)],
        capture_output=True,
        text=True,
        check=False,
    )
    if tested.returncode != 0:
        return False, tested.stderr or tested.stdout
    without_password = subprocess.run(
        [executable, "t", "-p", "-y", str(output)],
        capture_output=True,
        text=True,
        check=False,
    )
    if without_password.returncode == 0:
        return False, "archive does not require a password"
    listed = subprocess.run(
        [executable, "l", "-slt", f"-p{password}", str(output)],
        capture_output=True,
        text=True,
        check=False,
    )
    members = {
        line.split(" = ", 1)[1].replace("\\", "/")
        for line in listed.stdout.splitlines()
        if line.startswith("Path = ")
    }
    expected = set(spec["expected_members"])
    return listed.returncode == 0 and expected <= members, f"members={sorted(members)}"


def _pdf_text(run: Path, fixture: Path, spec: dict[str, Any]) -> tuple[float, str]:
    output = run / spec["output"]
    expected = fixture / spec["expected"]
    actual_text = _pdf_extracted_text(output)
    expected_text = _pdf_extracted_text(expected)
    score = fuzz.ratio(actual_text, expected_text) / 100
    return score, f"actual_chars={len(actual_text)}, similarity={score:.4f}"


def _pdf_extracted_text(path: Path) -> str:
    with fitz.open(path) as document:
        return "".join(page.get_text() for page in document).strip()


def _csv_file(run: Path, fixture: Path, spec: dict[str, Any]) -> tuple[bool, str]:
    output = run / spec["output"]
    expected = fixture / spec["expected"]
    with output.open(newline="", encoding="utf-8-sig") as actual_file:
        actual = list(csv.reader(actual_file))
    with expected.open(newline="", encoding="utf-8-sig") as expected_file:
        gold = list(csv.reader(expected_file))
    return actual == gold, f"actual_rows={len(actual)}, expected_rows={len(gold)}"


def _pdf_set(run: Path, fixture: Path, spec: dict[str, Any]) -> tuple[bool, str]:
    output = run / spec["output"]
    expected = set(spec["expected_members"])
    actual = {path.name for path in output.glob("*.pdf")}
    for path in output.glob("*.pdf"):
        PdfReader(path)
    return actual == expected, f"actual={sorted(actual)}, expected={sorted(expected)}"


def _docx_records(run: Path, fixture: Path, spec: dict[str, Any]) -> tuple[bool, str]:
    document = Document(run / spec["output"])
    records = json.loads((fixture / spec["expected"]).read_text(encoding="utf-8"))
    text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + ["\t".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
    ).casefold()
    values = [str(value).casefold() for record in records for value in record.values() if value]
    return all(value in text for value in values), f"checked_values={len(values)}"


def _yaml_records(run: Path, fixture: Path, spec: dict[str, Any]) -> tuple[bool, str]:
    output = run / spec["output"]
    document = yaml.safe_load(output.read_text(encoding="utf-8"))
    failed = []
    for rule in spec["rules"]:
        value: Any = document
        for key in rule["path"].split("."):
            value = value[key]
        if value != rule["equals"]:
            failed.append(rule["path"])
    return not failed, f"failed_rules={failed}"


def _terminal_command_and_workspace(
    run: Path, fixture: Path, spec: dict[str, Any]
) -> tuple[bool, str]:
    transcript = (run / spec["terminal_transcript"]).read_text(encoding="utf-8")
    expected = (run / spec["workspace"]).resolve()
    profile = (run / spec["vscode_profile"]).resolve()
    workspaces = _vscode_profile_workspaces(profile)
    profile_in_use = _vscode_profile_in_use(profile)
    case_terminal = _case_terminal_at_workspace(run, spec["workspace"])
    success = (
        spec["command"].casefold() in transcript.casefold()
        and expected in workspaces
        and profile_in_use
        and case_terminal
    )
    return success, (
        f"profile_workspaces={sorted(str(path) for path in workspaces)}, "
        f"profile_in_use={profile_in_use}, "
        f"case_terminal_at_workspace={case_terminal}"
    )


def _vscode_profile_workspaces(profile: Path) -> set[Path]:
    storage = profile / "User" / "workspaceStorage"
    workspaces = set()
    for state_path in storage.glob("*/workspace.json"):
        state = json.loads(state_path.read_text(encoding="utf-8"))
        folder = state.get("folder")
        if not isinstance(folder, str) or urlsplit(folder).scheme != "file":
            continue
        uri_path = unquote(urlsplit(folder).path)
        if re.match(r"^/[A-Za-z]:/", uri_path):
            uri_path = uri_path[1:]
        workspaces.add(Path(uri_path).resolve())
    return workspaces


def _vscode_profile_in_use(profile: Path) -> bool:
    expected = profile.resolve()
    for process in psutil.process_iter(["name", "cmdline"]):
        name = (process.info.get("name") or "").casefold()
        if name not in {"code", "code.exe"}:
            continue
        command = process.info.get("cmdline") or []
        for index, argument in enumerate(command):
            if argument.startswith("--user-data-dir="):
                value = argument.split("=", 1)[1]
            elif argument == "--user-data-dir" and index + 1 < len(command):
                value = command[index + 1]
            else:
                continue
            if Path(value).resolve() == expected:
                return True
    return False


def _case_terminal_at_workspace(run: Path, workspace: str) -> bool:
    metadata = json.loads((run / "case-processes.json").read_text(encoding="utf-8"))
    terminal_pid = metadata["terminal_pid"]
    if not psutil.pid_exists(terminal_pid):
        return False
    terminal_directory = Path(psutil.Process(terminal_pid).cwd()).resolve()
    return terminal_directory == (run / workspace).resolve()


def _process_and_transcript(
    run: Path, fixture: Path, spec: dict[str, Any]
) -> tuple[bool, str]:
    metadata = json.loads((run / spec["process_metadata"]).read_text(encoding="utf-8"))
    transcript = (run / spec["terminal_transcript"]).read_text(encoding="utf-8")
    writer_pids = set(metadata["pids"])
    running = [pid for pid in writer_pids if psutil.pid_exists(pid)]
    termination_command = _terminates_case_process(transcript, writer_pids)
    success = (
        spec["required_terminal_evidence"].casefold() in transcript.casefold()
        and termination_command
        and not running
    )
    return success, (
        f"termination_command={termination_command}, "
        f"running_case_pids={running}"
    )


def _terminates_case_process(transcript: str, pids: set[int]) -> bool:
    if not re.search(r"\b(?:stop-process|taskkill(?:\.exe)?)\b", transcript, re.IGNORECASE):
        return False
    ids = {
        int(value)
        for value in re.findall(r"(?:-id|/pid)\s+(\d+)", transcript, re.IGNORECASE)
    }
    if ids & pids:
        return True
    return bool(
        re.search(
            r"(?:-name\s+soffice|/im\s+soffice(?:\.exe)?)",
            transcript,
            re.IGNORECASE,
        )
    )


def _speedtest_export_time(
    run: Path, fixture: Path, spec: dict[str, Any]
) -> tuple[bool, str]:
    files = list((run / spec["output_directory"]).glob(spec["file_glob"]))
    if not files:
        return False, f"matching_files={len(files)}"
    output = max(files, key=lambda path: path.stat().st_mtime_ns)
    with output.open(newline="", encoding="utf-8-sig") as file:
        rows = list(csv.reader(file))
    value = rows[spec["timestamp_row"] - 1][spec["timestamp_column"] - 1]
    match = re.search(r"(\d{1,2}):(\d{2})", value)
    if match is None:
        return False, f"timestamp={value!r}"
    now = datetime.now()
    observed_minutes = int(match.group(1)) * 60 + int(match.group(2))
    now_minutes = now.hour * 60 + now.minute
    difference = abs(observed_minutes - now_minutes)
    difference = min(difference, 24 * 60 - difference)
    return difference < spec["max_age_minutes"], f"age_minutes={difference}"


def _structured_text(
    run: Path, fixture: Path, spec: dict[str, Any]
) -> tuple[bool, str]:
    text = (run / spec["output"]).read_text(encoding="utf-8")
    missing = [item for item in spec["required_text"] if item not in text]
    timestamp_lines = sum(
        bool(re.search(r"\b(?:[01]\d|2[0-3]):[0-5]\d:(?:[0-5]\d|60)\b", line))
        for line in text.splitlines()
    )
    success = not missing and timestamp_lines == spec["timestamp_line_count"]
    return success, f"missing={missing}, timestamp_lines={timestamp_lines}"


CHECKS: dict[str, Check] = {
    "text": _text,
    "same_file": _same_file,
    "hidden": _hidden,
    "all_files_moved": _all_files_moved,
    "image_size": _image_size,
    "docx_fonts": _docx_fonts,
    "workbook_values": _workbook_values,
    "png_names": _png_names,
    "seven_zip": _seven_zip,
    "pdf_text": _pdf_text,
    "csv": _csv_file,
    "pdf_set": _pdf_set,
    "docx_records": _docx_records,
    "yaml_records": _yaml_records,
    "terminal_command_and_workspace": _terminal_command_and_workspace,
    "process_and_transcript": _process_and_transcript,
    "speedtest_export_time": _speedtest_export_time,
    "structured_text": _structured_text,
}
